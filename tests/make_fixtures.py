"""Generate real OOXML fixtures for tests/test-officedoc.mjs.

Uses python-docx, openpyxl, python-pptx and Pillow (all bundled with the DSH
primary runtime). Run:

    <python> kbpro/tests/make_fixtures.py

Creates kbpro/tests/fixtures/{sample.docx,sample.xlsx,sample.pptx,marker.png}.
"""

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Font
from PIL import Image
from pptx import Presentation
from pptx.util import Inches as PptxInches

FIXTURES = Path(__file__).resolve().parent / "fixtures"
FIXTURES.mkdir(parents=True, exist_ok=True)

INJECTED = '5 < 10 & "quoted"'


def make_png(path: Path) -> None:
    image = Image.new("RGB", (120, 80), (32, 96, 192))
    for x in range(120):
        for y in range(80):
            if (x + y) % 16 == 0:
                image.putpixel((x, y), (240, 200, 40))
    image.save(path, format="PNG")


def add_hyperlink(paragraph, url: str, text: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)

    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def build_docx(png_path: Path) -> None:
    doc = Document()

    doc.add_heading("Quarterly Report", level=1)
    doc.add_heading("Financial Highlights", level=2)

    para = doc.add_paragraph("This is ")
    bold = para.add_run("BoldMarker")
    bold.bold = True
    para.add_run(" and ")
    italic = para.add_run("ItalicMarker")
    italic.italic = True
    para.add_run(" in one sentence.")

    colored = doc.add_paragraph()
    colored_run = colored.add_run("ColoredMarker")
    colored_run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    colored_run.font.size = Pt(14)

    escaped = doc.add_paragraph(f"Escaping check: {INJECTED}")

    doc.add_paragraph("BulletAlpha", style="List Bullet")
    doc.add_paragraph("BulletBeta", style="List Bullet")

    table = doc.add_table(rows=2, cols=3)
    headers = ["Region", "Revenue", "Growth"]
    for index, value in enumerate(headers):
        table.cell(0, index).text = value
    table.cell(1, 0).text = "TableCellMarker"
    table.cell(1, 1).text = "1250000"
    table.cell(1, 2).text = "12.5%"

    link_para = doc.add_paragraph("See the ")
    add_hyperlink(link_para, "https://example.com/report", "ExampleLink")
    link_para.add_run(" for details.")

    doc.add_paragraph("Chart below:")
    doc.add_picture(str(png_path), width=Inches(1.5))

    doc.core_properties.title = "Quarterly Report Fixture"
    doc.core_properties.author = "Fixture Author"
    doc.core_properties.subject = "Testing"

    doc.save(FIXTURES / "sample.docx")


def build_xlsx() -> None:
    wb = Workbook()

    data = wb.active
    data.title = "Data"
    data["A1"] = "Product"
    data["B1"] = "Qty"
    data["C1"] = "Price"
    data["D1"] = "Date"
    data["A2"] = "Widget"
    data["B2"] = 42
    data["C2"] = 19.5
    data["D2"] = datetime(2024, 3, 15)
    data["A3"] = "Gadget"
    data["B3"] = 7
    data["C3"] = 3.25
    data["D3"] = datetime(2024, 4, 1)
    data["E1"] = "Share"
    data["E2"] = 0.25
    data["E2"].number_format = "0%"
    data["A5"] = "NotesCell"
    data["B5"] = "inline text"

    summary = wb.create_sheet("Summary")
    summary["B3"] = "SparseCell"
    summary["D5"] = 1234.5
    summary["A1"] = "SummaryHeader"
    summary["A1"].font = Font(bold=True)

    wb.save(FIXTURES / "sample.xlsx")


def build_pptx() -> None:
    prs = Presentation()

    slides = [
        (
            "QuarterlyReportSlide",
            ["SlideOneBullet", "Second bullet on slide one", "Third bullet"],
        ),
        (
            "Operations Update",
            ["SlideTwoBullet", "Headcount stable", "Costs down 4%"],
        ),
        (
            "Next Steps",
            ["SlideThreeBullet", "Ship the beta", "Review in May"],
        ),
    ]

    for title, bullets in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        body = slide.placeholders[1].text_frame
        body.text = bullets[0]
        for bullet in bullets[1:]:
            para = body.add_paragraph()
            para.text = bullet
            para.level = 0

    prs.core_properties.title = "Presentation Fixture"
    prs.core_properties.author = "Fixture Author"

    prs.save(FIXTURES / "sample.pptx")


def main() -> None:
    png_path = FIXTURES / "marker.png"
    make_png(png_path)
    build_docx(png_path)
    build_xlsx()
    build_pptx()

    for name in ("sample.docx", "sample.xlsx", "sample.pptx", "marker.png"):
        path = FIXTURES / name
        print(f"{name}: {path.stat().st_size} bytes")


if __name__ == "__main__":
    main()
